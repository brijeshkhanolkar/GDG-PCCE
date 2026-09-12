import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

doc = docx.Document()

# Set standard margins
for section in doc.sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

# Color Palette
COLOR_GOLD = RGBColor(201, 162, 75)       # #C9A24B
COLOR_DARK = RGBColor(11, 11, 12)          # #0B0B0C
COLOR_MUTED = RGBColor(107, 107, 110)      # #6B6B6E
COLOR_RED = RGBColor(140, 59, 52)          # #8C3B34
COLOR_STEEL = RGBColor(142, 142, 147)      # #8E8E93

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

# Title Page / Header Block
title_p = doc.add_paragraph()
title_p.paragraph_format.space_before = Pt(10)
title_p.paragraph_format.space_after = Pt(2)
run_title = title_p.add_run("COURTFLIGHT")
run_title.font.name = "Arial"
run_title.font.size = Pt(26)
run_title.font.bold = True
run_title.font.color.rgb = COLOR_GOLD

sub_p = doc.add_paragraph()
sub_p.paragraph_format.space_after = Pt(14)
run_sub = sub_p.add_run("Unsupervised K-Means Clustering & Telemetry Engine")
run_sub.font.name = "Arial"
run_sub.font.size = Pt(16)
run_sub.font.bold = True
run_sub.font.color.rgb = COLOR_DARK

desc_p = doc.add_paragraph()
desc_p.paragraph_format.space_after = Pt(20)
run_desc = desc_p.add_run("Technical Architecture, Mathematical Foundations, Feature Vector Space, and Empirical Live Traces")
run_desc.font.name = "Arial"
run_desc.font.size = Pt(11)
run_desc.font.italic = True
run_desc.font.color.rgb = COLOR_MUTED

# Horizontal Divider Line
divider_table = doc.add_table(rows=1, cols=1)
divider_table.alignment = WD_TABLE_ALIGNMENT.CENTER
cell = divider_table.cell(0, 0)
set_cell_background(cell, "C9A24B")
cell.height = Pt(3)
cell.width = Inches(6.9)
doc.add_paragraph().paragraph_format.space_after = Pt(14)

# Section 1: Executive Summary
h1 = doc.add_heading(level=1)
h1_run = h1.add_run("1. Executive Summary & Core Objective")
h1_run.font.color.rgb = COLOR_GOLD
h1_run.font.size = Pt(15)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(8)
p.paragraph_format.line_spacing = 1.15
p.add_run(
    "CourtFlight upgrades Indian court case management from arbitrary manual estimates into a scientific, "
    "unsupervised machine learning engine. In conventional legal systems, pendency timelines are treated as black boxes. "
    "CourtFlight formulates every court case into an empirical 33-dimensional coordinate point and clusters it against "
    "the full dataset of 1,550 historical dockets across Indian High Courts."
)

p2 = doc.add_paragraph()
p2.paragraph_format.space_after = Pt(14)
p2.paragraph_format.line_spacing = 1.15
p2.add_run(
    "100% of all active and historical cases in CourtFlight are now classified via a precomputed, 10-cluster K-Means++ model. "
    "This document details the exact mathematical vectorization, Euclidean distance minimization, empirical percentile calculations, "
    "and a full execution trace demonstrating how docket NYR-00001 is mapped to Cluster #4."
)

# Section 2: Mathematical Architecture
h2 = doc.add_heading(level=1)
h2_run = h2.add_run("2. 5-Step Clustering Pipeline")
h2_run.font.color.rgb = COLOR_GOLD
h2_run.font.size = Pt(15)

steps = [
    ("Step 1: 33-Dimensional Vector Space Projection", 
     "Each case is vectorized into 33 numerical coordinates:\n"
     "• 6 Dimensions: One-hot encoded case type (Civil, Commercial, Criminal, Family, Property, Service)\n"
     "• 12 Dimensions: One-hot encoded filing court (Telangana, Bombay, Calcutta, Delhi, Madras, etc.)\n"
     "• 12 Dimensions: One-hot encoded state jurisdiction\n"
     "• 1 Dimension: Min-Max scaled case age in years: norm(age) = (age - minAge) / (maxAge - minAge)\n"
     "• 1 Dimension: Min-Max scaled adjournments count: norm(adjs) = (adjs - minAdj) / (maxAdj - minAdj)\n"
     "• 1 Dimension: Ordinal procedural stage index scaled 0.0 (Filed) to 1.0 (Disposed)"),
     
    ("Step 2: K-Means++ Optimization Across 1,550 Cases",
     "The engine was trained across 1,550 cases with k=10 clusters. To avoid local minima, 5 random restarts "
     "with k-means++ centroid initialization were executed. The clustering solution with the absolute lowest "
     "Within-Cluster Sum of Squares (WCSS = 2396.12) was precomputed and cached into cluster-stats.json."),
     
    ("Step 3: Real-Time Vector Inference & Euclidean Distance",
     "When an active case x is inspected, its Euclidean distance to every cluster centroid μ_k is computed:\n"
     "d(x, μ_k) = sqrt( sum_{i=1}^{33} (x_i - μ_{k,i})^2 )"),
     
    ("Step 4: Optimal Centroid Assignment",
     "The case is assigned to whichever cluster minimizes the Euclidean metric:\n"
     "k* = argmin_{k} d(x, μ_k)\n"
     "Because each of the 10 clusters contains between 18 and 97 disposed cases (average 46), every cluster is mature "
     "and trusted (method: 'cluster')."),
     
    ("Step 5: Empirical Percentile Extraction & Precedent Retrieval",
     "Instead of arbitrary formulas, the ETA window is derived strictly from historical cases in cluster k*:\n"
     "• p25: 25th percentile of years-to-disposal (accelerated scenario)\n"
     "• Median: 50th percentile of years-to-disposal (most probable trajectory)\n"
     "• p75: 75th percentile of years-to-disposal (extended litigation window)\n"
     "The 5 actual precedent cases with disposal dates closest to the median are sampled and displayed.")
]

for title, content in steps:
    p_s = doc.add_paragraph()
    p_s.paragraph_format.space_before = Pt(6)
    p_s.paragraph_format.space_after = Pt(2)
    r_s = p_s.add_run(title)
    r_s.font.bold = True
    r_s.font.size = Pt(11)
    r_s.font.color.rgb = COLOR_DARK

    p_c = doc.add_paragraph()
    p_c.paragraph_format.space_after = Pt(8)
    p_c.paragraph_format.line_spacing = 1.15
    r_c = p_c.add_run(content)
    r_c.font.size = Pt(10)
    r_c.font.color.rgb = COLOR_DARK

# Section 3: Table of 10 Precomputed Clusters
h3 = doc.add_heading(level=1)
h3_run = h3.add_run("3. The 10 Precomputed Cluster Baselines")
h3_run.font.color.rgb = COLOR_GOLD
h3_run.font.size = Pt(15)

doc.add_paragraph(
    "Below is the complete architectural baseline generated by scripts/build-clusters.js and saved to data/cluster-stats.json:"
).paragraph_format.space_after = Pt(8)

clusters_data = [
    ("Cluster #0", "277", "71", "3.7 yrs", "2.8 – 4.7 yrs", "Civil, Commercial, Property"),
    ("Cluster #1", "215", "76", "3.7 yrs", "2.8 – 4.5 yrs", "Civil, Commercial, Service"),
    ("Cluster #2", "147", "44", "4.0 yrs", "3.1 – 4.8 yrs", "Civil, Criminal, Service"),
    ("Cluster #3", "87", "32", "4.0 yrs", "3.0 – 5.1 yrs", "Civil, Property, Criminal"),
    ("Cluster #4", "68", "25", "3.9 yrs", "3.1 – 4.6 yrs", "Criminal (High Turbulence)"),
    ("Cluster #5", "130", "38", "4.2 yrs", "3.2 – 5.2 yrs", "Civil, Family, Property"),
    ("Cluster #6", "95", "18", "3.3 yrs", "2.4 – 4.1 yrs", "Civil, Family, Commercial"),
    ("Cluster #7", "343", "97", "4.4 yrs", "3.3 – 5.4 yrs", "Civil, Property, Service"),
    ("Cluster #8", "107", "29", "3.9 yrs", "3.0 – 4.9 yrs", "Property, Civil, Criminal"),
    ("Cluster #9", "81", "27", "4.4 yrs", "3.4 – 5.3 yrs", "Criminal (Standard Trajectory)"),
]

table = doc.add_table(rows=len(clusters_data) + 1, cols=6)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
headers = ["Cluster ID", "Total Cases", "Disposed", "Median ETA", "ETA Window (p25–p75)", "Dominant Litigation Vectors"]

# Format Header Row
hdr_row = table.rows[0]
for i, name in enumerate(headers):
    cell = hdr_row.cells[i]
    set_cell_background(cell, "0B0B0C")
    set_cell_margins(cell, top=140, bottom=140, left=140, right=140)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i in [1, 2, 3, 4] else WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(name)
    run.font.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = COLOR_GOLD

# Format Data Rows
for row_idx, data in enumerate(clusters_data):
    row = table.rows[row_idx + 1]
    bg_color = "F9F8F6" if row_idx % 2 == 0 else "FFFFFF"
    for col_idx, text in enumerate(data):
        cell = row.cells[col_idx]
        set_cell_background(cell, bg_color)
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx in [1, 2, 3, 4] else WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(text)
        run.font.size = Pt(8.5)
        run.font.color.rgb = COLOR_DARK
        if col_idx == 0:
            run.font.bold = True
        if col_idx == 3:
            run.font.bold = True
            run.font.color.rgb = COLOR_GOLD

doc.add_paragraph().paragraph_format.space_after = Pt(14)

# Section 4: Live Execution Trace of NYR-00001
h4 = doc.add_heading(level=1)
h4_run = h4.add_run("4. Live Execution Trace: Case NYR-00001")
h4_run.font.color.rgb = COLOR_GOLD
h4_run.font.size = Pt(15)

doc.add_paragraph(
    "To observe how clusterisation occurs in practice, consider docket NYR-00001 querying the engine in real time:"
).paragraph_format.space_after = Pt(6)

trace_box = doc.add_table(rows=1, cols=1)
trace_box.alignment = WD_TABLE_ALIGNMENT.CENTER
t_cell = trace_box.cell(0, 0)
set_cell_background(t_cell, "F4F3EF")
set_cell_margins(t_cell, top=140, bottom=140, left=160, right=160)

tp = t_cell.paragraphs[0]
tp.paragraph_format.line_spacing = 1.2
tp.add_run("DOCKET METRIC SPECIFICATION:\n").bold = True
tp.add_run("• Case ID: NYR-00001\n")
tp.add_run("• Case Type: Criminal | Filing Court: Telangana High Court | Jurisdiction: Telangana\n")
tp.add_run("• Current Stage: Arguments (Ordinal value = 0.667)\n")
tp.add_run("• Adjournments: 6 (Normalized = 0.667) | Filing Date: 2017-08-12 (Normalized Age = 0.841)\n\n")

tp.add_run("EUCLIDEAN DISTANCE COMPUTED TO ALL 10 CENTROIDS:\n").bold = True
tp.add_run("  Cluster #0 (Civil / Commercial / Property) : 1.6716\n")
tp.add_run("  Cluster #1 (Civil / Commercial / Service)  : 1.9493\n")
tp.add_run("  Cluster #2 (Civil / Criminal / Service)    : 2.2022\n")
tp.add_run("  Cluster #3 (Civil / Property / Criminal)   : 2.2143\n")
tp.add_run("  Cluster #4 (Criminal Dominant)             : 1.2682  <=== MINIMUM DISTANCE (SELECTED)\n")
tp.add_run("  Cluster #5 (Civil / Family / Property)     : 2.2016\n")
tp.add_run("  Cluster #6 (Civil / Family / Commercial)   : 2.2803\n")
tp.add_run("  Cluster #7 (Civil / Property / Service)    : 2.3284\n")
tp.add_run("  Cluster #8 (Property / Civil / Criminal)   : 2.2259\n")
tp.add_run("  Cluster #9 (Criminal Secondary)            : 2.0141\n\n")

tp.add_run("ENGINE INFERENCE RESULTS:\n").bold = True
tp.add_run("• Method Assigned: 'cluster' (Precomputed K-Means inference)\n")
tp.add_run("• Matched Cluster: Cluster #4 (68 cases in cluster)\n")
tp.add_run("• Empirical Percentiles: p25 = 3.1 yrs, Median = 3.9 yrs, p75 = 4.6 yrs\n")
tp.add_run("• Adjournment Benchmark: Case has 6 adjournments vs Cluster Avg of 3.7 => High Risk\n")
tp.add_run("• Precedent Cases Sampled from Cluster #4:\n")
tp.add_run("   1. NYR-••199 — Criminal, Telangana HC (3.9 yrs to disposal, 6 adjs)\n")
tp.add_run("   2. NYR-••015 — Criminal, Telangana HC (3.8 yrs to disposal, 9 adjs)\n")
tp.add_run("   3. NYR-••292 — Criminal, Rajasthan HC (3.9 yrs to disposal, 6 adjs)\n")
tp.add_run("   4. NYR-••761 — Criminal, Punjab & Haryana HC (3.9 yrs to disposal, 5 adjs)\n")
tp.add_run("   5. NYR-••814 — Criminal, Calcutta HC (3.8 yrs to disposal, 3 adjs)\n")

doc.add_paragraph().paragraph_format.space_after = Pt(14)

# Section 5: Verification Matrix
h5 = doc.add_heading(level=1)
h5_run = h5.add_run("5. Multi-Case Validation Across Featured Dockets")
h5_run.font.color.rgb = COLOR_GOLD
h5_run.font.size = Pt(15)

doc.add_paragraph(
    "To confirm that all cases are properly clustering and producing realistic windows across diverse courts and litigation vectors, "
    "the table below details the featured demo dockets:"
).paragraph_format.space_after = Pt(8)

validation_data = [
    ("NYR-00001", "Criminal", "Telangana HC", "Cluster #4", "68", "3.9 yrs", "3.1 – 4.6 yrs", "High"),
    ("NYR-00006", "Civil", "Bombay HC", "Cluster #1", "215", "3.7 yrs", "2.8 – 4.5 yrs", "Medium"),
    ("NYR-00009", "Property", "Calcutta HC", "Cluster #1", "215", "3.7 yrs", "2.8 – 4.5 yrs", "Medium"),
    ("NYR-00015", "Criminal", "Telangana HC", "Cluster #4", "68", "3.9 yrs", "3.1 – 4.6 yrs", "High"),
    ("NYR-00017", "Family", "Calcutta HC", "Cluster #1", "215", "3.7 yrs", "2.8 – 4.5 yrs", "Medium"),
    ("NYR-00013", "Property", "Bombay HC", "Cluster #7", "343", "4.4 yrs", "3.3 – 5.4 yrs", "High"),
    ("NYR-00002", "Family", "Bombay HC", "Cluster #7", "343", "4.4 yrs", "3.3 – 5.4 yrs", "High"),
]

v_table = doc.add_table(rows=len(validation_data) + 1, cols=8)
v_table.alignment = WD_TABLE_ALIGNMENT.CENTER
v_headers = ["Docket", "Type", "Filing Court", "Cluster", "Cluster Size", "Median ETA", "ETA Window", "Risk"]

hdr_row = v_table.rows[0]
for i, name in enumerate(v_headers):
    cell = hdr_row.cells[i]
    set_cell_background(cell, "0B0B0C")
    set_cell_margins(cell, top=140, bottom=140, left=100, right=100)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i in [3, 4, 5, 6, 7] else WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(name)
    run.font.bold = True
    run.font.size = Pt(8.5)
    run.font.color.rgb = COLOR_GOLD

for row_idx, data in enumerate(validation_data):
    row = v_table.rows[row_idx + 1]
    bg_color = "F9F8F6" if row_idx % 2 == 0 else "FFFFFF"
    for col_idx, text in enumerate(data):
        cell = row.cells[col_idx]
        set_cell_background(cell, bg_color)
        set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx in [3, 4, 5, 6, 7] else WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(text)
        run.font.size = Pt(8.5)
        run.font.color.rgb = COLOR_DARK
        if col_idx == 0:
            run.font.bold = True
        if col_idx == 7 and text == "High":
            run.font.color.rgb = COLOR_RED
            run.font.bold = True

doc.add_paragraph().paragraph_format.space_after = Pt(14)

# Section 6: API & UI Telemetry Integration
h6 = doc.add_heading(level=1)
h6_run = h6.add_run("6. Production Integration & Verification Endpoints")
h6_run.font.color.rgb = COLOR_GOLD
h6_run.font.size = Pt(15)

api_p = doc.add_paragraph()
api_p.paragraph_format.line_spacing = 1.15
api_p.add_run("The clusterisation engine operates seamlessly across two primary application interfaces:\n\n")
api_p.add_run("1. REST API Endpoint: GET /api/predict?id={DOCKET_ID}\n").bold = True
api_p.add_run(
    "Returns complete clustering telemetry including method ('cluster'), clusterId (e.g. 4), matchedClusterSize (68), "
    "empirical etaRangeYears ({p25, median, p75}), etaDateRange, clusterAvgAdjournments, and the 5 nearest precedent cases.\n\n"
)
api_p.add_run("2. Interactive Frontend UI: /case/{DOCKET_ID}\n").bold = True
api_p.add_run(
    "Renders the Section D 'Cluster Intelligence' component, showcasing the active K-Means Cluster badge, "
    "the total cluster size, the estimated resolution window, and nearest precedent cases with percentage procedural match."
)

# Footer Note
doc.add_paragraph().paragraph_format.space_after = Pt(20)
ft_p = doc.add_paragraph()
ft_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
ft_r = ft_p.add_run("CourtFlight Jurisdictional Analytics Telemetry // Generated by Antigravity Core")
ft_r.font.size = Pt(8.5)
ft_r.font.color.rgb = COLOR_MUTED

output_file = "D:/GDGPCCE/CourtFlight_Clustering_Architecture_Report.docx"
doc.save(output_file)
print(f"Document successfully created at: {output_file}")
